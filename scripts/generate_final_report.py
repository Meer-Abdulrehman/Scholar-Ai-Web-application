from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt, Inches
import os

project_title = "Scholar.AI — Student Performance & Career Recommendation System"
team = ["Rana Gulzaib", "Abdulrehman"]
authors = ", ".join(team)
model_name = "Scholar.AI"
output_path = "docs/scholar_ai_project_report_final.docx"
logo_path = "docs/university_logo.png"

if not os.path.exists('docs'):
    os.makedirs('docs')

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(12)

# Optional logo - center if available
if os.path.exists(logo_path):
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    try:
        p_logo.add_run().add_picture(logo_path, width=Inches(1.6))
    except Exception:
        # fallback: skip image if insertion fails
        pass

# Title page
p = doc.add_paragraph()
p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
r = p.add_run(project_title + '\n')
r.bold = True
r.font.size = Pt(26)

p2 = doc.add_paragraph()
p2.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
p2.add_run('Final Project Report\n').font.size = Pt(12)

p3 = doc.add_paragraph()
p3.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
p3.add_run(f'Team: {authors}\n')
p3.add_run(f'Model: {model_name}\n')
p3.add_run('University of South Asia, Lahore\n')
p3.add_run('Date: May 31, 2026\n')

doc.add_page_break()

# Define sections with medium-length paragraph content (Scholar.AI-specific)
sections = [
    ('Abstract',
     'Scholar.AI is an integrated system that predicts student performance and provides personalized, track-aware career recommendations. By combining supervised machine learning, explainability via SHAP, semantic career search, and an LLM-based advising agent, the platform enables timely interventions and informed decisions for students and advisors.'),

    ('Problem Statement',
     'Many secondary-school students do not receive individualized guidance that ties their subject strengths to realistic career pathways. Existing solutions often separate grade prediction from career counselling or produce generic advice without considering subject-specific tracks such as Pre-Engineering or Pre-Medical. Scholar.AI addresses this gap by delivering explainable predictions and track-specific career recommendations that reflect students\' merit and subject affinities.'),

    ('Introduction',
     'Scholar.AI was developed to help students, teachers, and career counselors identify academic strengths and translate them into actionable career pathways. The system integrates a React frontend, FastAPI backend, ML models for grade forecasting, a career recommender, and an LLM-driven advisor to provide concise, evidence-based recommendations.'),

    ('Literature Review',
     'Research in educational data mining demonstrates that past performance, attendance, and subject-specific scores are strong predictors of future outcomes. Explainable AI techniques have been used to make model outputs actionable. Scholar.AI builds on these foundations and combines them with modern semantic search and LLM affordances to better match students to careers.'),

    ('Data Collection',
     'The dataset is derived from public student performance records (UCI / Kaggle variants), merged and cleaned to form a representative sample. Key features include subject grades (math, science, english, computer, biology), study hours, attendance, and prior term results. Data was processed and saved under the project\'s data/processed folder.'),

    ('Exploratory Data Analysis (EDA)',
     'EDA revealed clear associations between prior grades, study effort, and final outcomes. Subject-level analysis highlights how track assignments (for example, stronger biology favoring Pre-Medical) inform career suitability. Visualizations and correlation analyses guided feature engineering and model selection.'),

    ('Feature Engineering',
     'Features such as previous-term averages, subject z-scores, and study-efficiency ratios were created to improve model sensitivity. Categorical variables were encoded consistently and continuous features were scaled so that models could learn robust, generalizable patterns.'),

    ('Modeling',
     'Regression and ensemble methods (Linear Regression, Random Forest) were trained to forecast final grades. A separate Random Forest classifier and deterministic rules were implemented for career recommendations, especially to ensure medically-appropriate outcomes for Pre-Medical students.'),

    ('Model Training Details',
     'Models were trained using an 80/20 train/test split and tuned via grid search for hyperparameters. Evaluation metrics included RMSE and R² for regression and accuracy/F1 for classification. Experiments were tracked using MLflow to capture parameters, metrics, and artifacts.'),

    ('Explainability (SHAP)',
     'SHAP was used to attribute feature importance at the individual prediction level, enabling tailored advice such as raising attendance or focusing study hours on weaker subjects. These explanations are surfaced in the frontend as charts and short textual recommendations.'),

    ('Career Recommendation System',
     'Scholar.AI uses a hybrid approach combining deterministic merit rules, semantic vector search over career descriptions, and LLM refinement. For Pre-Medical students, Biology mark ranges map to tiered medical careers (support roles to MBBS/BDS), ensuring recommendations remain aligned with merit and track constraints.'),

    ('Agent & MCP Server',
     'A lightweight agent coordinates prediction, explanation, and recommendation tools. The MCP server exposes these capabilities to other clients and provides an LLM-driven conversational layer for follow-up advice and personalized next steps.'),

    ('Vector DB & Semantic Search',
     'Career descriptions and profiles are embedded and stored in a vector-enabled database. Semantic search enables matching beyond simple score thresholds, helping discover niche or allied-health careers compatible with student strengths.'),

    ('TTS Integration',
     'Advice text is convertible to speech via a TTS service, improving accessibility. The frontend can play generated audio or download the advice clip for offline use.'),

    ('Frontend Architecture',
     'The frontend is implemented using React, Vite, and Tailwind for a responsive UI. Core components include the StudentForm, ResultCard, CareerChart, and AIAdviceBox, which together provide an integrated analysis view for users.'),

    ('Backend Architecture',
     'The backend uses FastAPI to expose prediction, recommendation, explanation, and TTS endpoints. Services are modularized under backend/services and models and artifacts are stored in backend/models for reproducibility.'),

    ('Deployment',
     'The application is deployable as containerized services. Environment variables manage secrets (DB URIs, API keys). Frontend hosting on Vercel or static hosts with a proxied API is supported; alternative deployments include Render or a VM for the backend.'),

    ('Security Considerations',
     'Secrets are stored out-of-repo and access to vector databases is restricted. Model artifacts should be validated before loading, and CORS should limit frontend origins for production deployments.'),

    ('Testing & Validation',
     'Unit tests cover core service logic while integration tests exercise API endpoints. Live end-to-end tests validate that frontend inputs yield coherent analysis, SHAP plots, and career recommendations.'),

    ('Limitations & Future Work',
     'Current limitations include dependence on public student datasets and the need for a larger labeled career dataset to fully data-drive recommendations. Future work includes fairness audits, multilingual support, and richer career taxonomies.'),

    ('Conclusion',
     'Scholar.AI demonstrates an effective, explainable approach to linking student performance with actionable career guidance. Its modular design supports iterative improvements and dataset expansion.'),

    ('References',
     'UCI Student Performance Dataset; scikit-learn; SHAP; LangChain; MongoDB Atlas Vector Search; Uplift Orator TTS; MLflow.'),

    ('Appendix A — Data Dictionary',
     'Subject grade fields: math, science, english, computer, biology (0–100). Additional fields include study_hours, attendance, prev_grade, and extracurricular flags.'),

    ('Appendix B — Model Artifacts',
     'Relevant artifacts are in backend/models (random_forest.pkl, linear_regression.pkl, rf_classifier.pkl, scaler.pkl) and scripts/train_career_recommender.py for career model training.'),

    ('Appendix C — API Examples',
     'POST /analyze with a student JSON returns prediction, SHAP, career recommendations, and AI advice. Example: curl -X POST http://localhost:8000/analyze -H "Content-Type: application/json" -d "{\"study_hours\":5,...}"')
]

# Build document: headings with medium paragraphs
for heading, paragraph in sections:
    doc.add_heading(heading, level=1)
    p = doc.add_paragraph(paragraph)
    p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    p.space_after = Pt(6)
    doc.add_page_break()

# Save
if os.path.exists(output_path):
    os.remove(output_path)

doc.save(output_path)
print(f"Generated final report: {output_path}")
