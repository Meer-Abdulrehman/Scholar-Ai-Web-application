from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

# Project-specific variables
project_title = "Scholar.AI — Student Performance & Career Recommendation System"
team = ["Abdulrehman (B-28721)", "Rana Muhammad Gulzaib (B-28506)"]
authors = ", ".join(team)
model_name = "scholar.ai"
output_path = "docs/scholar_ai_project_report.docx"

# Create document
doc = Document()

# Set default font
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

# Title Page
title = doc.add_paragraph()
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
run = title.add_run(project_title + "\n")
run.bold = True
run.font.size = Pt(24)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
run = subtitle.add_run("Project Report\n\n")
run.font.size = Pt(14)

meta = doc.add_paragraph()
meta.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
meta.add_run(f"Team: {authors}\n")
meta.add_run(f"Model: {model_name}\n")
meta.add_run("University of South Asia, Lahore\n")
meta.add_run("Date: May 31, 2026\n")

# Page break after title
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import nsmap

doc.add_page_break()

# Table of Contents placeholder
doc.add_heading('Table of Contents', level=1)
for i in range(1, 16):
    p = doc.add_paragraph(f"{i}. Section placeholder ....................... {i+1}")

# Ensure TOC page break
doc.add_page_break()

# Content sections to reach ~27 pages — we add explicit page breaks between sections
sections = [
    ("Executive Summary", "A short summary of the project, goals, and outcomes."),
    ("Introduction", "Problem statement, scope, and objectives."),
    ("Literature Review", "Related work and comparison to existing student performance systems."),
    ("Data Collection", "Datasets used, source (Kaggle student-mat/student-por), preprocessing steps."),
    ("Exploratory Data Analysis", "Key findings: correlations, distributions, plots."),
    ("Feature Engineering", "Selected features and transformations (scaling, encoding)."),
    ("Modeling", "Models trained: Linear Regression, Random Forest, Classifier for pass/fail."),
    ("Model Training Details", "Hyperparameters, cross-validation, MLflow tracking."),
    ("Explainability", "SHAP analysis and interpretation for individual predictions."),
    ("Career Recommendation System", "Fallback profiles, LLM integration, deterministic pre-medical rules."),
    ("Agent & MCP Server", "LangChain agent, tools, Groq LLM, MCP integration."),
    ("Vector DB & Semantic Search", "MongoDB Atlas Vector Search, career embeddings."),
    ("TTS Integration", "Uplift Orator Studio for audio advice."),
    ("Frontend Architecture", "React + Vite + Tailwind, key components and UI flows."),
    ("Backend Architecture", "FastAPI endpoints, services, models, DB connections."),
    ("Deployment", "Docker, docker-compose, recommended deployment steps."),
    ("Security Considerations", "Model unpickling, secrets handling, API keys."),
    ("Testing & Validation", "How to test endpoints and UI, unit/integration tests."),
    ("Limitations & Future Work", "Data limitations, improvements, dataset expansion."),
    ("Conclusion", "Summary and final remarks."),
    ("References", "Citations and data source references."),
    ("Appendix A — Data Dictionary", "Column descriptions and value ranges."),
    ("Appendix B — Model Artifacts", "Saved models and how to load them."),
    ("Appendix C — Sample API Calls", "Example curl commands and expected responses."),
]

# Helper to add a section with filler paragraphs
lorem = (
    "This section documents the project details. The Scholar.AI system uses student-provided features "
    "to predict grades, explain them using SHAP, and recommend suitable career paths. The system combines "
    "classical ML (scikit-learn), model explainability (SHAP), semantic search (MongoDB Atlas Vector Search), "
    "and an agentic LLM integration via LangChain."
)

for idx, (heading, intro) in enumerate(sections):
    doc.add_heading(heading, level=1)
    doc.add_paragraph(intro)
    # Add multiple paragraphs to fill the page roughly
    for p in range(8):
        doc.add_paragraph(lorem)
    # Add a page break to control page count
    doc.add_page_break()

# Add a short appendix table in final page
from docx.shared import Inches

doc.add_heading('Appendix: Quick Reference', level=1)
para = doc.add_paragraph()
para.add_run('API Endpoint Reference:').bold = True

table = doc.add_table(rows=1, cols=3)
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Endpoint'
hdr_cells[1].text = 'Method'
hdr_cells[2].text = 'Description'

rows = [
    ('/analyze', 'POST', 'All-in-one predict + recommend + explain + advice'),
    ('/predict', 'POST', 'Grade prediction only'),
    ('/recommend', 'POST', 'Career recommendations'),
    ('/explain', 'POST', 'SHAP explanation'),
]
for r in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = r[0]
    row_cells[1].text = r[1]
    row_cells[2].text = r[2]

# Save document
import os
os.makedirs(os.path.dirname(output_path), exist_ok=True)
doc.save(output_path)
print(f"Saved report to {output_path}")
